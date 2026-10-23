#!/usr/bin/env python3
"""Read published v1.2.0 exports; project with GeographicLib, never PyProj.

Usage: python audit.py ASSET_DIR TRANSVERSE_MERCATOR_PROJ OUTPUT_DIR
The asset directory's parent must contain release.json and release-tag.json.
"""

import csv
import hashlib
import io
import json
import math
import re
import sqlite3
import statistics
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
import shapefile

THRESHOLD_M = 0.02
PARAMETERS = ["-e", "6378137", "1/298.257222101", "-l", "19", "-k", "0.9993", "-p", "9"]


def require(condition, description):
    if not condition:
        raise ValueError(description)


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def indexed(rows):
    result = {}
    for row in rows:
        key = (str(row["object_id"]), str(row["measurement_id"]))
        require(all(key) and key not in result, f"Empty/duplicate key: {key}")
        result[key] = row
    return result


def project(executable, coordinates):
    run = subprocess.run(
        [str(executable), *PARAMETERS],
        input="".join(f"{lat:.17g} {lon:.17g}\n" for lat, lon in coordinates),
        text=True,
        capture_output=True,
        check=True,
    )
    lines = run.stdout.splitlines()
    require(len(lines) == len(coordinates), "Projection output cardinality")
    result = []
    for line in lines:
        east, north, convergence, scale = map(float, line.split())
        require(all(map(math.isfinite, (east, north, convergence, scale))), line)
        result.append((north - 5300000, east + 500000))
    return result


def wkt_values(text):
    match = re.fullmatch(r"POINT\(([^ ]+) ([^ ]+)\)", text)
    require(match is not None, f"Invalid WKT: {text}")
    return tuple(map(float, match.groups()))


def excel_rows(sheet):
    rows = sheet.iter_rows(values_only=True)
    for row in rows:
        if row[0] == "ID obiektu":
            headers = row
            break
    else:
        raise ValueError("Missing Excel header")
    return [dict(zip(headers, row, strict=True)) for row in rows if row[0] is not None]


def describe(rows):
    values = sorted(row["distance_m"] for row in rows)
    require(bool(values), "Empty representation")
    return {
        "rows": len(rows),
        "over_0_02_m": sum(value > THRESHOLD_M for value in values),
        "max_distance_m": values[-1],
        "median_distance_m": statistics.median(values),
        "p95_distance_m": values[math.ceil(len(values) * 0.95) - 1],
        "max_abs_delta_x_m": max(abs(row["delta_x_m"]) for row in rows),
        "max_abs_delta_y_m": max(abs(row["delta_y_m"]) for row in rows),
    }


def main():
    assets, executable, output = map(Path, sys.argv[1:])
    output.mkdir(parents=True, exist_ok=True)
    version = subprocess.run(
        [str(executable), "--version"], text=True, capture_output=True, check=True
    ).stdout.strip()
    require(version.endswith("GeographicLib version 2.7"), "GeographicLib version")
    provenance = assets.parent
    release = json.loads((provenance / "release.json").read_text())
    tag = json.loads((provenance / "release-tag.json").read_text())
    manifest = []
    for asset in release["assets"]:
        path = assets / asset["name"]
        digest = sha256(path)
        require(path.stat().st_size == asset["size"], f"Size: {path.name}")
        require("sha256:" + digest == asset["digest"], f"Digest: {path.name}")
        manifest.append(
            {
                "name": path.name,
                "size": path.stat().st_size,
                "sha256": digest,
                "url": asset["browser_download_url"],
            }
        )
    require(len(manifest) == 9, "Expected complete 9-file release")
    stats = json.loads((assets / "statistics.json").read_text())
    require(stats["release"]["version"] == release["tag_name"] == "v1.2.0", "Version")
    require(stats["release"]["source_commit"] == tag["object"]["sha"], "Tag/commit")

    controls = project(executable, [(0, 19), (49, 19), (49, 18), (49, 20)])
    require(controls[0] == (-5300000, 500000), "Equator/origin and offsets")
    require(controls[1][1] == 500000, "Central meridian/easting")
    require(abs(controls[2][0] - controls[3][0]) < 1e-8, "Symmetric northing")
    require(abs(controls[2][1] + controls[3][1] - 1000000) < 1e-8, "Symmetric easting")

    with zipfile.ZipFile(assets / "katalog.sqlite.zip") as archive:
        require(archive.namelist() == ["katalog.sqlite"], "SQLite ZIP contents")
        database_bytes = archive.read("katalog.sqlite")
    with tempfile.TemporaryDirectory() as temp:
        database_path = Path(temp) / "katalog.sqlite"
        database_path.write_bytes(database_bytes)
        with sqlite3.connect(database_path.as_uri() + "?mode=ro", uri=True) as connection:
            connection.row_factory = sqlite3.Row
            require(
                connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok",
                "SQLite integrity",
            )
            require(
                not connection.execute("PRAGMA foreign_key_check").fetchall(), "SQLite foreign keys"
            )
            measurements = [
                dict(row)
                for row in connection.execute(
                    "SELECT *, id AS measurement_id FROM measurements ORDER BY object_id,id"
                )
            ]
            objects = [dict(row) for row in connection.execute("SELECT * FROM objects ORDER BY id")]
            selections = {
                row["object_id"]: row["measurement_id"]
                for row in connection.execute(
                    "SELECT object_id,measurement_id FROM best_measurements"
                )
            }
    baseline = indexed(measurements)
    selected_keys = {(row["id"], row["best_measurement_id"]) for row in objects}
    require(len(selected_keys) == len(objects) == len(selections), "Selection cardinality")
    require(all((key, value) in selected_keys for key, value in selections.items()), "Selections")
    require(selected_keys <= baseline.keys(), "Selection references")
    require(len(baseline) == stats["counts"]["measurements"], "Measurement count")
    require(len(objects) == stats["counts"]["objects"], "Object count")
    representations = []
    coverage = {}
    joins = {}

    def add(
        name,
        rows,
        *,
        all_measurements=False,
        native_xy=True,
        degree_tolerance=1e-12,
        xy_tolerance=1e-9,
    ):
        records = indexed(rows)
        expected = baseline.keys() if all_measurements else selected_keys
        require(records.keys() == expected, f"Missing/extra keys: {name}")
        max_degree_difference = 0
        max_xy_difference = 0
        for key, row in records.items():
            original = baseline[key]
            lat, lon = float(row["lat"]), float(row["lon"])
            require(math.isfinite(lat) and math.isfinite(lon), f"Nonfinite WGS84: {name}/{key}")
            require(-90 <= lat <= 90 and -180 <= lon <= 180, f"WGS84 range: {name}/{key}")
            d = max(abs(lat - original["lat"]), abs(lon - original["lon"]))
            max_degree_difference = max(max_degree_difference, d)
            require(d <= degree_tolerance, f"WGS84 differs: {name}/{key}: {d}")
            x = float(row["x_1992"] if native_xy else original["x_1992"])
            y = float(row["y_1992"] if native_xy else original["y_1992"])
            require(math.isfinite(x) and math.isfinite(y), f"Nonfinite XY: {name}/{key}")
            dxy = max(abs(x - original["x_1992"]), abs(y - original["y_1992"]))
            max_xy_difference = max(max_xy_difference, dxy)
            require(dxy <= xy_tolerance, f"XY differs: {name}/{key}: {dxy}")
            representations.append(
                {
                    "format": name,
                    "object_id": key[0],
                    "measurement_id": key[1],
                    "source": original["source"],
                    "selected": key in selected_keys,
                    "native_xy": native_xy,
                    "lat": lat,
                    "lon": lon,
                    "x_1992": x,
                    "y_1992": y,
                }
            )
        coverage[name] = len(records)
        joins[name] = {
            "max_wgs84_difference_deg": max_degree_difference,
            "max_xy_difference_m": max_xy_difference,
            "degree_tolerance": degree_tolerance,
            "xy_tolerance_m": xy_tolerance,
            "native_xy": native_xy,
        }

    add("SQLite measurements", measurements, all_measurements=True)
    object_rows = []
    for row in objects:
        require(
            wkt_values(row["best_geom_wgs84"]) == (row["best_lon"], row["best_lat"]),
            "Object WKT WGS84 axes",
        )
        require(
            wkt_values(row["best_geom_1992"]) == (row["best_y_1992"], row["best_x_1992"]),
            "Object WKT PL1992 axes",
        )
        object_rows.append(
            {
                "object_id": row["id"],
                "measurement_id": row["best_measurement_id"],
                **{field: row["best_" + field] for field in ("lat", "lon", "x_1992", "y_1992")},
            }
        )
    for row in measurements:
        require(
            wkt_values(row["geom_wgs84"]) == (row["lon"], row["lat"]), "Measurement WKT WGS84 axes"
        )
        require(
            wkt_values(row["geom_1992"]) == (row["y_1992"], row["x_1992"]),
            "Measurement WKT PL1992 axes",
        )
    add("SQLite objects", object_rows)
    with (assets / "best-measurements.csv").open(newline="", encoding="utf-8") as file:
        add("CSV", list(csv.DictReader(file)), degree_tolerance=5.01e-11)

    geojson = json.loads((assets / "best-measurements.geojson").read_text())
    geometry_rows = []
    properties = []
    require(geojson["type"] == "FeatureCollection", "GeoJSON collection")
    for feature in geojson["features"]:
        prop = feature["properties"]
        geometry = feature["geometry"]
        require(geometry["type"] == "Point" and len(geometry["coordinates"]) == 2, "GeoJSON point")
        lon, lat = geometry["coordinates"]
        require((lon, lat) == (prop["lon"], prop["lat"]), "GeoJSON lon/lat axes")
        properties.append(prop)
        geometry_rows.append({**prop, "lat": lat, "lon": lon})
    add("GeoJSON properties", properties)
    add("GeoJSON geometry", geometry_rows, native_xy=False)

    gpx = []
    root = ET.parse(assets / "best-measurements.gpx").getroot()
    for waypoint in root.findall("{*}wpt"):
        desc = waypoint.findtext("{*}desc", "")
        match = re.search(r"; object ([^;]+); measurement ([^;]+);", desc)
        require(match is not None, "GPX identity")
        gpx.append(
            {
                "object_id": match[1],
                "measurement_id": match[2],
                "lat": waypoint.attrib["lat"],
                "lon": waypoint.attrib["lon"],
            }
        )
    add("GPX", gpx, native_xy=False, degree_tolerance=5.01e-11)

    with zipfile.ZipFile(assets / "best-measurements.shp.zip") as archive:
        prefix = "best-measurements"
        prj = archive.read(prefix + ".prj").decode("ascii")
        for parameter, expected in {
            "False_Easting": 500000,
            "False_Northing": -5300000,
            "Central_Meridian": 19,
            "Scale_Factor": 0.9993,
            "Latitude_Of_Origin": 0,
        }.items():
            match = re.search(r'PARAMETER\["' + parameter + r'",([-0-9.]+)\]', prj)
            require(match is not None and float(match[1]) == expected, "SHP CRS: " + parameter)
        require('SPHEROID["GRS_1980",6378137.0,298.257222101]' in prj, "SHP GRS80")
        require(
            'PROJECTION["Transverse_Mercator"]' in prj and 'UNIT["Meter",1.0]' in prj,
            "SHP projection/unit",
        )
        reader = shapefile.Reader(
            **{
                extension: io.BytesIO(archive.read(prefix + "." + extension))
                for extension in ("shp", "shx", "dbf")
            }
        )
        attributes, shapes = [], []
        for feature in reader.iterShapeRecords():
            row = feature.record.as_dict()
            require(
                feature.shape.shapeType == shapefile.POINT and len(feature.shape.points) == 1,
                "SHP Point",
            )
            east, north = feature.shape.points[0]
            row["measurement_id"] = row["meas_id"]
            attributes.append(row)
            shapes.append({**row, "x_1992": north, "y_1992": east})
        add("Shapefile attributes", attributes, degree_tolerance=5.01e-9, xy_tolerance=0.000501)
        add("Shapefile geometry", shapes, degree_tolerance=5.01e-9)

    workbook = openpyxl.load_workbook(assets / "katalog.xlsx", read_only=True, data_only=True)
    require(workbook.sheetnames == ["Best", "Full"], "Excel sheets")
    for name in workbook.sheetnames:
        rows = excel_rows(workbook[name])
        converted = []
        for row in rows:
            point = {
                "object_id": row["ID obiektu"],
                "measurement_id": row["ID pomiaru"],
                "lat": row["Szerokość WGS84 [°]"],
                "lon": row["Długość WGS84 [°]"],
            }
            if name == "Full":
                point.update(x_1992=row["X PL-1992 [m]"], y_1992=row["Y PL-1992 [m]"])
                selected = (point["object_id"], point["measurement_id"]) in selected_keys
                require(
                    row["Wybrany do eksportu"] == ("Tak" if selected else "Nie"), "Excel selection"
                )
            converted.append(point)
        add(
            "Excel " + name,
            converted,
            all_measurements=name == "Full",
            native_xy=name == "Full",
            degree_tolerance=1e-12 if name == "Full" else 5.01e-11,
        )
    workbook.close()

    projected = project(executable, [(row["lat"], row["lon"]) for row in representations])
    for row, (x, y) in zip(representations, projected, strict=True):
        row.update(
            reference_x=x, reference_y=y, delta_x_m=row["x_1992"] - x, delta_y_m=row["y_1992"] - y
        )
        row["distance_m"] = math.hypot(row["delta_x_m"], row["delta_y_m"])
        row["over_0_02_m"] = row["distance_m"] > THRESHOLD_M
    by_format = {
        name: describe([row for row in representations if row["format"] == name])
        for name in coverage
    }
    canonical = [row for row in representations if row["format"] == "SQLite measurements"]
    for row in canonical:
        key = (row["object_id"], row["measurement_id"])
        row["max_distance_all_formats_m"] = max(
            item["distance_m"]
            for item in representations
            if (item["object_id"], item["measurement_id"]) == key
        )
    fields = [name for name in representations[0] if name != "max_distance_all_formats_m"]
    for filename, rows, columns in [
        ("representations.csv", representations, fields),
        ("points.csv", canonical, [*fields, "max_distance_all_formats_m"]),
    ]:
        with (output / filename).open("w", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(
                file, fieldnames=columns, extrasaction="ignore", lineterminator="\n"
            )
            writer.writeheader()
            writer.writerows(rows)
    summary = {
        "release": stats["release"],
        "published_at": release["published_at"],
        "release_url": release["html_url"],
        "assets": manifest,
        "geographiclib": {
            "version": "2.7",
            "commit": "15cddd7eed9f7f61ead6dba2229d45b271dbdeda",
            "binary_sha256": sha256(executable),
            "parameters": PARAMETERS,
            "controls_xy": controls,
        },
        "threshold_m": THRESHOLD_M,
        "coverage": coverage,
        "cross_format": joins,
        "by_format": by_format,
        "all_measurements": describe(canonical),
        "selected_measurements": describe([row for row in canonical if row["selected"]]),
        "all_representations": describe(representations),
        "worst_measurements": sorted(canonical, key=lambda row: row["distance_m"], reverse=True)[
            :10
        ],
        "unexplained_differences": sum(row["over_0_02_m"] for row in representations),
        "by_source": {
            source: describe([row for row in canonical if row["source"] == source])
            for source in sorted({row["source"] for row in canonical})
        },
        "checks": {
            "asset_hashes": "PASS",
            "sqlite_integrity_and_fk": "PASS",
            "unique_and_complete_ids": "PASS",
            "geometry_axes_and_shp_crs": "PASS",
            "cross_format_values_with_declared_rounding": "PASS",
        },
    }
    (output / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    print(
        json.dumps(
            {
                key: summary[key]
                for key in (
                    "coverage",
                    "by_format",
                    "all_measurements",
                    "selected_measurements",
                    "all_representations",
                    "unexplained_differences",
                )
            },
            indent=2,
        )
    )
    return 1 if summary["unexplained_differences"] else 0


if __name__ == "__main__":
    sys.exit(main())
