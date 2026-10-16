import csv
import json
import subprocess
import sys
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
import yaml
from openpyxl import load_workbook

from gps_kataster_obiektow_tatr.coordinates import wgs84_to_1992
from gps_kataster_obiektow_tatr.excel_export import write_excel_export
from gps_kataster_obiektow_tatr.release_artifacts import build_release_artifacts

REPO_ROOT = Path(__file__).resolve().parents[1]
BUILD_RELEASE_SCRIPT = REPO_ROOT / "scripts" / "build_release_artifacts.py"
KSW_LAT = 49.23459299
KSW_LON = 19.87589498
GENERATED_AT = "2026-05-16T12:00:00Z"


def test_build_release_artifacts_writes_full_expected_set(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    sqlite_path = tmp_path / "build" / "katalog.sqlite"
    output_dir = tmp_path / "build" / "exports"
    _write_sample_data(data_dir)

    result = build_release_artifacts(
        data_dir=data_dir,
        sqlite_path=sqlite_path,
        output_dir=output_dir,
        generated_at=GENERATED_AT,
    )

    assert result.artifact_paths == (
        sqlite_path,
        output_dir / "best-measurements.geojson",
        output_dir / "best-measurements.csv",
        output_dir / "best-measurements.gpx",
        output_dir / "best-measurements.shp.zip",
        output_dir / "katalog.sqlite.zip",
        output_dir / "metadata.json",
        output_dir / "statistics.json",
        output_dir / "statistics.md",
        output_dir / "katalog.xlsx",
    )
    for artifact_path in result.artifact_paths:
        assert artifact_path.exists()

    with zipfile.ZipFile(output_dir / "katalog.sqlite.zip") as archive:
        assert archive.namelist() == ["katalog.sqlite"]

    assert result.export_result.metadata["counts"] == {
        "objects": 1,
        "caves": 1,
        "relations": 0,
        "measurements": 2,
        "validation_warnings": 0,
        "validation_errors": 0,
    }
    statistics = json.loads((output_dir / "statistics.json").read_text())
    assert statistics["selected_points"]["horizontal_accuracy"] == {"known": 1, "unknown": 0}


def test_build_release_artifacts_cli_is_local_dry_run(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    sqlite_path = tmp_path / "build" / "katalog.sqlite"
    output_dir = tmp_path / "build" / "exports"
    _write_sample_data(data_dir)

    result = subprocess.run(
        [
            sys.executable,
            str(BUILD_RELEASE_SCRIPT),
            "--data-dir",
            str(data_dir),
            "--sqlite-output",
            str(sqlite_path),
            "--output-dir",
            str(output_dir),
            "--generated-at",
            GENERATED_AT,
            "--release-version",
            "v1.2.3",
            "--source-commit",
            "a" * 40,
        ],
        cwd=REPO_ROOT,
        check=False,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    for filename in (
        "best-measurements.geojson",
        "best-measurements.csv",
        "best-measurements.gpx",
        "best-measurements.shp.zip",
        "katalog.sqlite.zip",
        "metadata.json",
        "statistics.json",
        "statistics.md",
        "katalog.xlsx",
    ):
        assert (output_dir / filename).exists()
    assert sqlite_path.exists()
    assert "release artifacts: 1 objects, 1 caves, 2 measurements" in result.stdout
    statistics = json.loads((output_dir / "statistics.json").read_text())
    assert statistics["release"] == {"version": "v1.2.3", "source_commit": "a" * 40}
    assert statistics["generated_at"] == GENERATED_AT


def test_statistics_count_selected_points_separately_from_history(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    _write_sample_data(data_dir)
    first = _valid_object()
    first["best_measurement"].update(
        mode="manual", measurement_id="m-001", reason="Confirmed older measurement"
    )
    first["measurements"][0].update(
        verification_status="zweryfikowany",
        horizontal_accuracy_m=0,
        verified_by="tester",
        verified_at="2026-05-16T10:00:00Z",
    )
    first["measurements"][1].update(verification_status="odrzucony", notes="Wrong position")
    _write_yaml(data_dir / "objects/KSW/KSW-0001.yml", first)
    second = _valid_object()
    second["id"] = "KSW-0002"
    second["external_refs"] = []
    second["category"] = "sztolnia"
    second["measurements"][1]["horizontal_accuracy_m"] = None
    _write_yaml(data_dir / "objects/KSW/KSW-0002.yml", second)
    cave = _valid_cave()
    cave["object_ids"].append("KSW-0002")
    _write_yaml(data_dir / "caves/C-0001.yml", cave)

    result = build_release_artifacts(
        data_dir=data_dir,
        sqlite_path=tmp_path / "db.sqlite",
        output_dir=tmp_path / "exports",
        generated_at=GENERATED_AT,
    )
    statistics = json.loads((result.export_result.output_dir / "statistics.json").read_text())
    assert statistics["release"] == {"version": None, "source_commit": None}
    assert statistics["counts"] == {"objects": 2, "caves": 1, "measurements": 4}
    assert statistics["selected_points"] == {
        "by_category": {"jaskinia_otwor": 1, "sztolnia": 1},
        "by_source": {"PIG": 1, "TPN": 1},
        "by_verification_status": {"nieweryfikowany": 1, "zweryfikowany": 1},
        "horizontal_accuracy": {"known": 1, "unknown": 1},
    }
    assert statistics["measurement_history"] == {
        "by_source": {"PIG": 2, "TPN": 2},
        "by_verification_status": {
            "nieweryfikowany": 2,
            "odrzucony": 1,
            "zweryfikowany": 1,
        },
    }
    report = (result.export_result.output_dir / "statistics.md").read_text()
    assert "| Otwory jaskiń | 1 |" in report
    assert "| Cave entrances | 1 |" in report
    assert "| Jaskynné vchody | 1 |" in report
    assert "| PIG | 1 | 2 |" in report
    assert "| Odrzucone | 0 | 1 |" in report


def test_statistics_are_deterministic_and_support_an_empty_catalog(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    for name in ("objects", "caves", "relations"):
        (data_dir / name).mkdir(parents=True)
    reports = []
    for name in ("first", "second"):
        output_dir = tmp_path / name
        build_release_artifacts(
            data_dir=data_dir,
            sqlite_path=output_dir / "db.sqlite",
            output_dir=output_dir,
            generated_at=GENERATED_AT,
        )
        reports.append(
            tuple((output_dir / f"statistics.{suffix}").read_bytes() for suffix in ("json", "md"))
        )
    assert reports[0] == reports[1]
    statistics = json.loads(reports[0][0])
    assert statistics["counts"] == {"objects": 0, "caves": 0, "measurements": 0}
    assert statistics["selected_points"]["horizontal_accuracy"] == {"known": 0, "unknown": 0}
    assert statistics["selected_points"]["by_source"] == {}


def test_excel_preserves_best_and_full_history_with_local_measurement_ids(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    _write_sample_data(data_dir)
    first = _valid_object()
    first["name_local"] = '=HYPERLINK("https://example.invalid", "name")'
    first["notes"] = "Original object notes"
    first["best_measurement"].update(mode="manual", measurement_id="m-001", reason="Older checked")
    first["measurements"][0]["horizontal_accuracy_m"] = 0
    first["measurements"][1].update(verification_status="odrzucony", notes="Rejected location")
    _write_yaml(data_dir / "objects/KSW/KSW-0001.yml", first)
    second = _valid_object()
    second["id"] = "KSW-0002"
    second["external_refs"] = []
    second["measurements"][1]["horizontal_accuracy_m"] = None
    _write_yaml(data_dir / "objects/KSW/KSW-0002.yml", second)
    cave = _valid_cave()
    cave["object_ids"].append("KSW-0002")
    _write_yaml(data_dir / "caves/C-0001.yml", cave)
    result = build_release_artifacts(
        data_dir=data_dir,
        sqlite_path=tmp_path / "db.sqlite",
        output_dir=tmp_path / "exports",
        generated_at=GENERATED_AT,
        release_version="v1.2.3",
        source_commit="a" * 40,
    )
    path = result.export_result.output_dir / "katalog.xlsx"
    assert path.is_file()
    book = load_workbook(path)
    assert book.sheetnames == ["Best", "Full"]
    best = _excel_records(book["Best"])
    full = _excel_records(book["Full"])
    assert len(best) == 2 and len(full) == 4
    assert [(r["ID obiektu"], r["ID pomiaru"]) for r in best] == [
        ("KSW-0001", "m-001"),
        ("KSW-0002", "m-002"),
    ]
    assert best[0]["Nazwa obiektu"] == first["name_local"]
    assert best[0]["Dokładność pozioma [m]"] == 0
    assert best[1]["Dokładność pozioma [m]"] is None
    assert best[0]["Szerokość WGS84 [°]"] == pytest.approx(KSW_LAT)
    assert best[0]["Data pomiaru"] == datetime(2026, 5, 15)
    assert [r["Wybrany do eksportu"] for r in full] == ["Tak", "Nie", "Nie", "Tak"]
    assert full[0]["Dokładność pozioma [m]"] == 0
    assert full[3]["Dokładność pozioma [m]"] is None
    assert full[1]["Weryfikacja"] == "odrzucony"
    assert full[1]["Uwagi pomiaru"] == "Rejected location"
    assert full[0]["Uwagi obiektu"] == first["notes"]
    assert full[0]["Nazwa jaskini"] == cave["name"]
    assert (
        json.loads(full[0]["Referencje obiektu (JSON)"])[0]["external_id"]
        == first["external_refs"][0]["external_id"]
    )
    assert json.loads(full[0]["Referencje jaskini (JSON)"])[0]["external_id"] == "1094"
    assert full[0]["Nadany prefiks"] == "KSW"
    assert full[0]["Utworzył obiekt"] == "dl"
    assert full[0]["Data źródła"] is None
    for sheet, count in ((book["Best"], 2), (book["Full"], 4)):
        assert sheet.freeze_panes == "C7"
        assert sheet.sheet_view.showGridLines is False
        table = next(iter(sheet.tables.values()))
        assert table.ref.endswith(str(count + 6))
        assert table.autoFilter is not None
        assert not sheet.merged_cells.ranges
        assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert "v1.2.3" in book.properties.description
    assert "a" * 40 in book.properties.description
    book.close()


@pytest.mark.parametrize(
    ("source_date", "expected"),
    [
        ("0001-01-01", "0001-01-01"),
        ("1895-06-15", "1895-06-15"),
        ("1899-12-31", "1899-12-31"),
        ("1900-01-01", datetime(1900, 1, 1)),
        ("1900-01-02", datetime(1900, 1, 2)),
        ("1900-02-28", datetime(1900, 2, 28)),
        ("1900-03-01", datetime(1900, 3, 1)),
        ("2024-02-29", datetime(2024, 2, 29)),
        ("9999-12-31", datetime(9999, 12, 31)),
    ],
)
def test_excel_preserves_dates_across_1900_boundary(
    tmp_path: Path, source_date: str, expected: str | datetime
) -> None:
    data_dir = tmp_path / "data"
    _write_sample_data(data_dir)
    obj = _valid_object()
    for measurement in obj["measurements"]:
        measurement.update(observed_date=source_date, source_date=source_date)
    _write_yaml(data_dir / "objects/KSW/KSW-0001.yml", obj)

    result = build_release_artifacts(
        data_dir=data_dir,
        sqlite_path=tmp_path / "db.sqlite",
        output_dir=tmp_path / "exports",
        generated_at=GENERATED_AT,
    )

    book = load_workbook(result.excel_path)
    try:
        best = _excel_records(book["Best"])
        full = _excel_records(book["Full"])
        assert [row["Data pomiaru"] for row in best] == [expected]
        assert [row["Data pomiaru"] for row in full] == [expected, expected]
        assert [row["Data źródła"] for row in full] == [expected, expected]
    finally:
        book.close()


def test_excel_empty_catalog_has_headers_without_fake_records(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    for name in ("objects", "caves", "relations"):
        (data_dir / name).mkdir(parents=True)
    build_release_artifacts(
        data_dir=data_dir,
        sqlite_path=tmp_path / "db.sqlite",
        output_dir=tmp_path / "exports",
        generated_at=GENERATED_AT,
    )
    path = tmp_path / "exports/katalog.xlsx"
    assert path.is_file()
    book = load_workbook(path)
    assert book.sheetnames == ["Best", "Full"]
    assert all(sheet.max_row == 6 for sheet in book)
    assert all(sheet["A6"].value == "ID obiektu" for sheet in book)
    book.close()


def test_excel_rejects_text_truncation_without_replacing_existing_workbook(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    _write_sample_data(data_dir)
    result = build_release_artifacts(
        data_dir=data_dir,
        sqlite_path=tmp_path / "db.sqlite",
        output_dir=tmp_path / "exports",
        generated_at=GENERATED_AT,
    )
    before = result.excel_path.read_bytes()
    with result.export_result.csv_path.open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    rows[0]["name_local"] = "x" * 32768
    with result.export_result.csv_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    with pytest.raises(ValueError, match="Cannot export Excel cell: Best"):
        write_excel_export(
            csv_path=result.export_result.csv_path,
            sqlite_path=result.sqlite_result.sqlite_path,
            output_path=result.excel_path,
            generated_at=GENERATED_AT,
            release_version=None,
            source_commit=None,
        )
    assert result.excel_path.read_bytes() == before
    assert not result.excel_path.with_suffix(".xlsx.tmp").exists()


def _excel_records(sheet: Any) -> list[dict[str, Any]]:
    headers = [cell.value for cell in sheet[6]]
    return [
        dict(zip(headers, row, strict=True)) for row in sheet.iter_rows(min_row=7, values_only=True)
    ]


def _write_sample_data(data_dir: Path) -> None:
    for name in ("objects", "caves", "relations"):
        (data_dir / name).mkdir(parents=True, exist_ok=True)
    _write_yaml(data_dir / "objects" / "KSW" / "KSW-0001.yml", _valid_object())
    _write_yaml(data_dir / "caves" / "C-0001.yml", _valid_cave())


def _valid_object() -> dict[str, Any]:
    return {
        "schema_version": 1,
        "id": "KSW-0001",
        "category": "jaskinia_otwor",
        "name_local": "Test Cave - main entrance",
        "cave_id": "C-0001",
        "id_assignment": {
            "method": "auto",
            "assigned_from_measurement_id": "m-001",
            "assigned_prefix": "KSW",
            "prefix_override_reason": None,
        },
        "external_refs": [
            {
                "system": "TPN",
                "ref_type": "source_globalid",
                "external_id": "{38626571-CAA6-4317-8900-D61A995020E9}",
                "scope": "object",
                "notes": "TPN point reference.",
            }
        ],
        "measurements": [
            _measurement("m-001", source="PIG", observed_date="2026-05-15"),
            _measurement("m-002", source="TPN", observed_date="2026-05-16"),
        ],
        "best_measurement": {
            "mode": "auto",
            "measurement_id": "m-002",
            "reason": None,
            "updated_at": "2026-05-16T10:30:00Z",
            "updated_by": "dl",
        },
        "attachments": [],
        "notes": None,
        "created_at": "2026-05-15T10:00:00Z",
        "created_by": "dl",
        "updated_at": "2026-05-16T10:30:00Z",
        "updated_by": "dl",
    }


def _valid_cave() -> dict[str, Any]:
    return {
        "schema_version": 1,
        "id": "C-0001",
        "name": "Test Cave",
        "system_name": None,
        "external_refs": [
            {
                "system": "PIG",
                "ref_type": "catalog_id",
                "external_id": "1094",
                "url": "https://jaskiniepolski.pgi.gov.pl/Details/Information/1094",
                "scope": "cave",
                "notes": "PIG catalog record identifier.",
            }
        ],
        "object_ids": ["KSW-0001"],
        "notes": None,
        "created_at": "2026-05-15T10:00:00Z",
        "created_by": "dl",
        "updated_at": "2026-05-16T10:30:00Z",
        "updated_by": "dl",
    }


def _measurement(
    measurement_id: str,
    *,
    source: str,
    observed_date: str,
) -> dict[str, Any]:
    pl_1992 = wgs84_to_1992(lat=KSW_LAT, lon=KSW_LON)
    return {
        "id": measurement_id,
        "lat": KSW_LAT,
        "lon": KSW_LON,
        "x_1992": pl_1992.x_1992,
        "y_1992": pl_1992.y_1992,
        "elevation_m": 1240.0,
        "elevation_datum": "unknown",
        "elevation_source": "source_record",
        "horizontal_accuracy_m": 5.0,
        "vertical_accuracy_m": 8.0,
        "source": source,
        "source_ref": f"{source}:{measurement_id}",
        "observed_at": None,
        "observed_date": observed_date,
        "source_date": None,
        "method": "source_record",
        "device": None,
        "tags": ["fixture"],
        "verification_status": "nieweryfikowany",
        "verified_by": None,
        "verified_at": None,
        "notes": None,
        "created_at": "2026-05-15T10:00:00Z",
        "created_by": "dl",
    }


def _write_yaml(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(deepcopy(data), sort_keys=False), encoding="utf-8")
