import hashlib
import importlib.util
import json
import os
import sqlite3
import subprocess
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from gps_kataster_obiektow_tatr.coordinates import wgs84_to_1992

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT = REPO_ROOT / "scripts" / "build_pr_data_preview.py"
HEAD = "a" * 40
BASE = "b" * 40
MERGE = "c" * 40
GENERATED_AT = "2026-10-20T12:34:56Z"
RELEASE_ASSETS = {
    "best-measurements.geojson",
    "best-measurements.csv",
    "best-measurements.gpx",
    "best-measurements.shp.zip",
    "katalog.sqlite.zip",
    "metadata.json",
    "statistics.json",
    "statistics.md",
    "katalog.xlsx",
}


def test_cli_builds_complete_release_preview_with_provenance_and_checksums(tmp_path: Path) -> None:
    data_dir = _sample_data(tmp_path)
    github_output = tmp_path / "github-output"
    result = _run(tmp_path, _arguments(data_dir), github_output=github_output)

    assert result.returncode == 0, result.stdout + result.stderr
    preview_root = tmp_path / "build" / "pr-data"
    exports = preview_root / "exports"
    assert {path.name for path in exports.iterdir()} == RELEASE_ASSETS | {
        "preview.json",
        "SHA256SUMS",
    }
    assert (preview_root / "katalog.sqlite").is_file()
    assert not (exports / "katalog.sqlite").exists()
    assert json.loads((exports / "preview.json").read_text()) == {
        "pr_number": 12,
        "head_commit": HEAD,
        "base_commit": BASE,
        "merge_commit": MERGE,
        "run_id": 12345,
        "run_attempt": 2,
        "generated_at": GENERATED_AT,
    }
    checksums = {}
    for line in (exports / "SHA256SUMS").read_text().splitlines():
        digest, filename = line.split("  ", 1)
        checksums[filename] = digest
    assert set(checksums) == RELEASE_ASSETS | {"preview.json"}
    for filename, digest in checksums.items():
        assert digest == hashlib.sha256((exports / filename).read_bytes()).hexdigest()

    statistics = json.loads((exports / "statistics.json").read_text())
    assert statistics["release"] == {
        "version": "pr-12-aaaaaaaaaaaa",
        "source_commit": MERGE,
    }
    assert statistics["generated_at"] == GENERATED_AT
    assert statistics["counts"] == {"objects": 1, "caves": 1, "measurements": 1}
    assert json.loads((exports / "metadata.json").read_text())["generated_at"] == GENERATED_AT
    with zipfile.ZipFile(exports / "katalog.sqlite.zip") as archive:
        assert archive.namelist() == ["katalog.sqlite"]
        assert archive.read("katalog.sqlite") == (preview_root / "katalog.sqlite").read_bytes()
    with sqlite3.connect(preview_root / "katalog.sqlite") as connection:
        metadata = dict(connection.execute("SELECT key, value FROM metadata"))
    assert metadata["generated_at"] == GENERATED_AT
    book = load_workbook(exports / "katalog.xlsx")
    assert book.sheetnames == ["Best", "Full"]
    assert book.properties.created == datetime(2026, 10, 20, 12, 34, 56)
    assert f"pr-12-aaaaaaaaaaaa; {MERGE}; {GENERATED_AT}" in book.properties.description
    for sheet in book:
        assert sheet.max_row == 7
        assert "pr-12-aaaaaaaaaaaa" in sheet["A3"].value

    assert github_output.read_text() == (
        f"artifact_name=pr-data-12-h{HEAD}-b{BASE}-m{MERGE}-r12345-a2\n"
    )
    assert "wrote:" in result.stdout


@pytest.mark.parametrize(
    ("option", "invalid"),
    [
        ("--pr-number", "0"),
        ("--pr-number", "-1"),
        ("--run-id", "0"),
        ("--run-attempt", "-1"),
        ("--head-commit", "A" * 40),
        ("--base-commit", "b" * 39),
        ("--merge-commit", "g" * 40),
        ("--generated-at", "not-a-timestamp"),
        ("--generated-at", "2026-10-20T12:34:56"),
    ],
)
def test_cli_rejects_invalid_provenance_before_creating_outputs(
    tmp_path: Path, option: str, invalid: str
) -> None:
    arguments = _arguments(tmp_path / "missing-data")
    arguments[arguments.index(option) + 1] = invalid
    github_output = tmp_path / "github-output"
    result = _run(tmp_path, arguments, github_output=github_output)

    assert result.returncode != 0
    assert option in result.stderr
    assert not (tmp_path / "build").exists()
    assert not github_output.exists()


@pytest.mark.parametrize(
    "option",
    [
        "--pr-number",
        "--head-commit",
        "--base-commit",
        "--merge-commit",
        "--run-id",
        "--run-attempt",
    ],
)
def test_cli_requires_complete_provenance(tmp_path: Path, option: str) -> None:
    arguments = _arguments(tmp_path / "missing-data")
    index = arguments.index(option)
    del arguments[index : index + 2]
    result = _run(tmp_path, arguments)

    assert result.returncode != 0
    assert option in result.stderr
    assert not (tmp_path / "build").exists()


def test_cli_refuses_stale_export_directory_without_removing_files(tmp_path: Path) -> None:
    exports = tmp_path / "build" / "pr-data" / "exports"
    exports.mkdir(parents=True)
    stale = exports / "stale.txt"
    stale.write_text("preserve this source")
    github_output = tmp_path / "github-output"

    result = _run(tmp_path, _arguments(_sample_data(tmp_path)), github_output=github_output)

    assert result.returncode != 0
    assert "not empty" in result.stderr
    assert stale.read_text() == "preserve this source"
    assert not (exports.parent / "katalog.sqlite").exists()
    assert not github_output.exists()


@pytest.mark.parametrize("corruption", ["missing", "extra", "directory", "symlink"])
def test_preview_rejects_incomplete_or_unexpected_release_files(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    corruption: str,
) -> None:
    spec = importlib.util.spec_from_file_location("pr_data_preview", SCRIPT)
    assert spec is not None and spec.loader is not None
    preview = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(preview)
    real_build = preview.build_release_artifacts

    def corrupt_build(**kwargs):
        result = real_build(**kwargs)
        exports = kwargs["output_dir"]
        if corruption == "missing":
            (exports / "katalog.xlsx").unlink()
        elif corruption == "extra":
            (exports / "leftover.csv").write_text("stale")
        elif corruption == "directory":
            (exports / "katalog.xlsx").unlink()
            (exports / "katalog.xlsx").mkdir()
        else:
            (exports / "katalog.xlsx").unlink()
            (exports / "katalog.xlsx").symlink_to(exports / "statistics.json")
        return result

    monkeypatch.setattr(preview, "build_release_artifacts", corrupt_build)
    github_output = tmp_path / "github-output"
    monkeypatch.setenv("GITHUB_OUTPUT", str(github_output))
    preview_root = tmp_path / "custom-preview"
    arguments = [*_arguments(_sample_data(tmp_path)), "--output-dir", str(preview_root)]

    assert preview.main(arguments) == 1
    assert "release asset" in capsys.readouterr().err
    assert not (preview_root / "exports" / "preview.json").exists()
    assert not (preview_root / "exports" / "SHA256SUMS").exists()
    assert not github_output.exists()


def _arguments(data_dir: Path) -> list[str]:
    return [
        "--data-dir",
        str(data_dir),
        "--pr-number",
        "12",
        "--head-commit",
        HEAD,
        "--base-commit",
        BASE,
        "--merge-commit",
        MERGE,
        "--run-id",
        "12345",
        "--run-attempt",
        "2",
        "--generated-at",
        GENERATED_AT,
    ]


def _run(
    cwd: Path, arguments: list[str], *, github_output: Path | None = None
) -> subprocess.CompletedProcess[str]:
    environment = dict(os.environ)
    environment.pop("GITHUB_OUTPUT", None)
    if github_output is not None:
        environment["GITHUB_OUTPUT"] = str(github_output)
    return subprocess.run(
        [sys.executable, str(SCRIPT), *arguments],
        cwd=cwd,
        env=environment,
        check=False,
        capture_output=True,
        text=True,
    )


def _sample_data(tmp_path: Path) -> Path:
    data_dir = tmp_path / "data"
    for kind in ("objects", "caves", "relations"):
        (data_dir / kind).mkdir(parents=True)
    fixture_dir = REPO_ROOT / "tests" / "fixtures"
    obj = yaml.safe_load((fixture_dir / "valid-object.yml").read_text())
    measurement = obj["measurements"][0]
    measurement.update(lat=49.23459299, lon=19.87589498, source_ref="fixture:m-001")
    projected = wgs84_to_1992(lat=measurement["lat"], lon=measurement["lon"])
    measurement.update(x_1992=projected.x_1992, y_1992=projected.y_1992)
    cave = yaml.safe_load((fixture_dir / "valid-cave.yml").read_text())
    cave["object_ids"] = [obj["id"]]
    (data_dir / "objects" / "KSW").mkdir()
    (data_dir / "objects" / "KSW" / "KSW-0001.yml").write_text(
        yaml.safe_dump(obj), encoding="utf-8"
    )
    (data_dir / "caves" / "C-0001.yml").write_text(yaml.safe_dump(cave), encoding="utf-8")
    return data_dir
